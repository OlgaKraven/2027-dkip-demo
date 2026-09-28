"""Generate authored production scenarios and their consistent input documents."""
from pathlib import Path
from datetime import date
from decimal import Decimal
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
root=Path(__file__).resolve().parents[1]
variants=[]
for i,line in enumerate((root/'content/domains.txt').read_text(encoding='utf8').splitlines(),1):
 domain,brand,product,process,rules,materials,operations=line.split('|')
 vid=f'V{i:02}'; folder=root/'content/variants'/vid;folder.mkdir(exist_ok=True)
 qty=[12,40,18,100,25,30][(i-1)%6]; resources=[]
 for entry in materials.split(';'):
  name,unit,norm,price=entry.split(',');resources.append(dict(name=name,kind='material',unit=unit,norm=float(norm),time=1,price=float(price)))
 for entry in operations.split(';'):
  name,hours,price=entry.split(',');resources.append(dict(name=name,kind='operation',unit='ч',norm=1,time=float(hours),price=float(price)))
 customers=[dict(id=f'{i:03}{j:06}',name=f'{n} «{brand.split("«")[1].rstrip("»")}-{j}»',inn='',addres=f'Учебный город, улица Производственная, {i}, офис {j}',phone=f'+7999{i:03}{j:04}',type='Покупатель') for j,n in enumerate(['Магазин','Торговый дом','Студия','Магазин подарков','Мастерская','Оптовый склад'],1)]
 common=f'{brand} выпускает продукцию «{product}» по спецификациям. {customers[0]["name"]} заказал {qty} ед. к 22.03.2027. Менеджер оформляет заказ, технолог задаёт состав, мастер планирует выпуск. Требуется связать документы в одной информационной системе и рассчитать нормативную стоимость заказа.'
 accounting='Нормы заданы на одну единицу готовой продукции. Материалы: количество выпуска × норма × цена. Операции: количество выпуска × часы × тариф за час. Цены действуют с 01.03.2027, НДС, наценка, доставка и дополнительные накладные расходы не учитываются. Промежуточные суммы не округлять, итог округлить до копеек.'
 texts=[
 f'По документам {brand} спроектируйте ER-диаграмму в 3НФ. Отразите заказчиков, продукцию «{product}», спецификации с материалами и операциями, цены и производственный заказ. Для каждого документа определите реквизиты и повторяющиеся строки. Укажите PK, FK, кратности; сохраните PDF.',
 f'Создайте базу для {brand} в выбранной СУБД по вашей схеме. Импортируйте все 6 записей Заказчики.json, сохранив ведущие нули кодов. Загрузите заказ {vid}-01, спецификацию и цены из Excel. Проверьте внешние ключи, сделайте выгрузку и восстановите её в пустую базу.',
 f'Создайте SQL-запрос нормативной стоимости заказа {vid}-01 предприятия {brand}: {qty} ед. продукции «{product}». {rules} {accounting} Покажите материалы и операции отдельными строками и общую сумму.',
 f'Создайте приложение {brand}: Windows Forms + MySQL либо WPF + PostgreSQL. Реализуйте вход по логину, паролю и пазлу, роли Администратор и Пользователь; третья подряд ошибка пароля или пазла блокирует запись. Администратор добавляет и изменяет пользователей, снимает блокировку. Проверьте пустые поля, дубли логинов и обработку ошибок. Используйте изображения и требования приложения к заданию 4.',
 f'Реализуйте API заметок сотрудников {brand}. Создайте users и notes со связью автора, внесите 5 записей из Заметки.json, сопоставив логины существующим пользователям. GET /notes или /api/notes возвращает id, title_user («title - login»), content без изменений, formatted_date (ДД.ММ.ГГГГ). Проверьте 200, пустой массив, 400 для неверного запроса и 500 при сбое БД. Сохраните коллекцию с исполняемыми проверками HTTP-кодов и JSON.',
 f'Документируйте API заметок {brand} в исходном шаблоне DOCX из приложения к заданию 6. Через Swagger UI выполните реальный запрос и перенесите его ответ в документ. В OpenAPI проверьте пути, параметры, поля и типы ответа. Укажите URL, реализованные методы, форматы, примеры 200, пустого результата, 400 и 500. Сохраните DOCX, OpenAPI JSON и тестовую коллекцию.'
 ]
 notes=[dict(title=t,login=l,content=c,created_date='2027-03-16') for t,l,c in [
 ('Заказ принят','manager',f'Заказ {vid}-01: {product}, {qty} ед.; срок 22.03.2027.'),
 ('Спецификация согласована','technologist',rules),('Материалы проверены','master',f'Подготовить: {", ".join(r["name"] for r in resources if r["kind"]=="material")}.') ,
 ('Порядок производства','master',process),('Расчёт заказа','manager','Использовать цены от 01.03.2027. Итог округлить до копеек.')]]
 def write_json(name,data): (folder/name).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
 write_json('Заказчики.json',customers);write_json('Заметки.json',notes)
 def book(name,rows,headers):
  wb=openpyxl.Workbook();ws=wb.active;ws.title='Данные'
  for row in rows:ws.append(row)
  for row in ws:
   for cell in row:
    cell.font=Font(name='Calibri',size=11)
    cell.alignment=Alignment(vertical='top',wrap_text=True)
    if isinstance(cell.value,date):cell.number_format='dd.mm.yyyy'
    elif isinstance(cell.value,float):cell.number_format='0' if cell.value.is_integer() else '0.###'
    if name=='Цены.xlsx' and cell.column==2 and cell.row>2:cell.number_format='0.00'
   ws.row_dimensions[row[0].row].height=42
  for n in headers:
   for cell in ws[n]:cell.fill=PatternFill('solid',fgColor='A61F32');cell.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF')
  for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=34 if col[0].column==1 else 25
  ws.freeze_panes=f'A{headers[-1]+1}';ws.sheet_view.showGridLines=False
  ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True
  ws.print_options.horizontalCentered=True
  wb.save(folder/name)
 book('Заказ покупателя.xlsx',[
 ['Авторский учебный вариант',vid,domain],['Номер заказа','Дата','Срок','Код заказчика'],[vid+'-01',date(2027,3,15),date(2027,3,22),customers[0]['id']],
 ['Изготовитель','Заказчик'],[brand,customers[0]['name']],['Продукция','Единица','Количество'],[product,'ед.',qty]], [1,2,4,6])
 book('Спецификация.xlsx',[
 ['Авторский учебный вариант',vid,domain],['Продукция',product,'На 1 единицу'],['Ресурс','Вид','Единица','Норма расхода','Часы на единицу']]+
 [[r['name'],r['kind'],r['unit'],r['norm'],r['time'] if r['kind']=='operation' else None] for r in resources], [1,3])
 book('Цены.xlsx', [['Авторский учебный вариант',vid,domain],['Ресурс','Цена, руб./ед.','Единица','Действует с']]+
 [[r['name'],r['price'],r['unit'],date(2027,3,1)] for r in resources], [1,2])
 book('Заказ на производство.xlsx',[
 ['Авторский учебный вариант',vid,domain],['Производственный заказ','Заказ покупателя','Дата запуска'],[vid+'-P01',vid+'-01',date(2027,3,16)],
 ['Продукция','Количество'],[product,qty],['Ресурс','Плановый расход','Единица']]+
 [[r['name'],float(Decimal(str(qty))*Decimal(str(r['norm']))*Decimal(str(r['time'] if r['kind']=='operation' else 1))),r['unit']] for r in resources],[1,2,4,6])
 tasks=[dict(text=t) for t in texts]
 tasks[5]['checks']=['Заполнен исходный шаблон DOCX.','Описаны все фактически реализованные пути и параметры.','Примеры взяты из ответа API выбранного предприятия.','Есть примеры 200, пустого массива, 400 и 500.','OpenAPI, Swagger UI, код и документ согласованы.']
 v=dict(id=vid,number=i,revision='domains-v2',domain=domain,brand=brand,product=product,common=common,process=process,rules=rules+' '+accounting,qty=qty,resources=resources,customers=customers,tasks=tasks,notes=notes)
 variants.append(v)
 md=f'# {vid} · {domain}\n\n{common}\n\n## Производственный процесс\n\n{process}\n\n## Правила учёта\n\n{rules} {accounting}\n\n'
 md+='Авторский учебный вариант по структуре базового КИМ 09.02.07-5-2027. Не официальный экзаменационный вариант. 6 заданий, 210 минут, 75 баллов. Все организации и данные учебные.\n\n'
 md+='Материалы и операции в спецификации различаются значениями material и operation. Время указано только для операций. Поле addres в JSON сохранено как имя входного поля. Пользователи manager, technologist и master — обычные пользователи, административную запись создайте отдельно.\n\n'
 md+='\n\n'.join(f'## Задание {n}\n\n{t}' for n,t in enumerate(texts,1))
 md+='\n\n## Результаты\n\nER PDF, SQL и выгрузка БД, запрос стоимости, проект приложения, проект API и коллекция проверок, DOCX по шаблону, OpenAPI JSON. Приложения к заданиям 4–6 находятся в каталоге «Приложения» архива.\n'
 (folder/'Задание.md').write_text(md,encoding='utf8')
(root/'content/variants.json').write_text(json.dumps(variants,ensure_ascii=False,indent=2),encoding='utf8')
print('Prepared 30 distinct domains, 120 workbooks, 60 JSON inputs and 30 six-task briefs.')
